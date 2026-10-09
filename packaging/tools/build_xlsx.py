"""Construit packaging_engine_mms_v2.xlsx : classeur d'origine + feuilles Packaging Calculator et Mixed Pack Helper."""
import sys, openpyxl
from copy import copy
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
bold = Font(name="Arial", bold=True)
normal = Font(name="Arial")
hdr_fill = PatternFill("solid", fgColor="DDEBF7")
inp_fill = PatternFill("solid", fgColor="FFFFCC")
thin = Side(style="thin", color="999999")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_hdr(ws, cells):
    for c in cells:
        ws[c].font = bold; ws[c].fill = hdr_fill; ws[c].border = box
        ws[c].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ---------- Produits : colonne Couchable? (Non par défaut) ----------
wsP = wb["Produits"]
hdrs = [str(c.value or "") for c in wsP[1]]
if not any(h.lower().startswith("couchable") for h in hdrs):
    col = wsP.max_column + 1
    ref = wsP.cell(1, 1)
    h = wsP.cell(1, col, "Couchable?")
    h.font = copy(ref.font); h.fill = copy(ref.fill); h.border = copy(ref.border); h.alignment = copy(ref.alignment)
    for r in range(2, wsP.max_row + 1):
        if wsP.cell(r, 1).value is not None:
            c = wsP.cell(r, col, "Non"); c.font = copy(wsP.cell(r, 4).font); c.border = copy(wsP.cell(r, 4).border); c.fill = inp_fill
    wsP.column_dimensions[get_column_letter(col)].width = 12
    if wsP.auto_filter.ref:
        wsP.auto_filter.ref = f"A1:{get_column_letter(col)}{wsP.max_row}"
    # validation Oui/Non
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    wsP.add_data_validation(dv); dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{wsP.max_row}")
    cmt = wsP.cell(1, col)
    cmt.comment = Comment("Oui = l'article peut être couché sur n'importe quelle face dans la caisse. Non (défaut) = il reste debout, seule la rotation à plat est permise (ex. sachets souples).", "Packaging Engine")

# ---------- Produits : colonne Caisses interdites ----------
hdrs = [str(c.value or "") for c in wsP[1]]
if not any(h.lower().startswith("caisses interdites") for h in hdrs):
    col = wsP.max_column + 1
    ref = wsP.cell(1, 1)
    h = wsP.cell(1, col, "Caisses interdites")
    h.font = copy(ref.font); h.fill = copy(ref.fill); h.border = copy(ref.border); h.alignment = copy(ref.alignment)
    h.comment = Comment("Noms exacts de la feuille Caisses, séparés par « ; » (ex. Caisse Super Small;Caisse V2). Ces caisses sont exclues pour toute commande contenant ce produit. Vide = aucune restriction.", "Packaging Engine")
    for r in range(2, wsP.max_row + 1):
        if wsP.cell(r, 1).value is not None:
            c = wsP.cell(r, col); c.font = copy(wsP.cell(r, 4).font); c.border = copy(wsP.cell(r, 4).border); c.fill = inp_fill
            if str(wsP.cell(r, 1).value).strip() == "9004":
                c.value = "Caisse Super Small"   # barème entrepôt : 1 sachet -> V2
    wsP.column_dimensions[get_column_letter(col)].width = 22
    if wsP.auto_filter.ref:
        wsP.auto_filter.ref = f"A1:{get_column_letter(col)}{wsP.max_row}"

# ---------- Packaging Calculator ----------
if "Packaging Calculator" in wb.sheetnames: del wb["Packaging Calculator"]
ws = wb.create_sheet("Packaging Calculator")
ws["A1"] = "Marge vide All Seasons"; ws["B1"] = 0.05
ws["A2"] = "Marge vide Summer";      ws["B2"] = 0.30
ws["A3"] = "Jeu de sécurité (cm)";   ws["B3"] = 0.5
for c in ("A1", "A2", "A3"): ws[c].font = bold
for c in ("B1", "B2", "B3"):
    ws[c].fill = inp_fill; ws[c].border = box; ws[c].font = Font(name="Arial", color="0000FF")
ws["B1"].number_format = "0%"; ws["B2"].number_format = "0%"; ws["B3"].number_format = "0.0"
ws["C1"] = "← % du volume de la caisse laissé vide (marge en haut). Modifiable."
ws["C2"] = "← plus large en été (gel / conservation). Modifiable."
ws["C3"] = "← cm retranchés aux 3 dimensions utiles de chaque caisse (articles trop justes). Modifiable."
for c in ("C1", "C2", "C3"): ws[c].font = Font(name="Arial", italic=True, color="666666")
ws["A4"] = "Item"; ws["B4"] = "Quantity"; ws["C4"] = "Suggested Package All Seasons"; ws["D4"] = "Suggested Package Summer"; ws["E4"] = "Info"
style_hdr(ws, ["A4", "B4", "C4", "D4", "E4"])
# lignes d'exemple (SKU de la feuille Produits)
examples = [("9004", 1), ("9004", 10), ("9124", 3), ("9130", 2), ("9081", 2), ("9150", 5), ("9491", 1)]
for i, (sku, q) in enumerate(examples, start=5):
    ws.cell(i, 1, sku).font = normal; ws.cell(i, 2, q).font = normal
    for c in range(1, 6): ws.cell(i, c).border = box
ws["E1"] = "Mode d'emploi : SKU (ou nom exact du produit) en colonne A, quantité en B, puis bouton « Remplir »."
ws["E2"] = "C, D, E sont remplies par la macro (format qté:caisse, séparateur « _ »)."
ws["E1"].font = Font(name="Arial", italic=True, color="666666"); ws["E2"].font = Font(name="Arial", italic=True, color="666666")
ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 11
ws.column_dimensions["C"].width = 36; ws.column_dimensions["D"].width = 36; ws.column_dimensions["E"].width = 60
ws.freeze_panes = "A5"

# ---------- Mixed Pack Helper ----------
if "Mixed Pack Helper" in wb.sheetnames: del wb["Mixed Pack Helper"]
ws = wb.create_sheet("Mixed Pack Helper")
hdrs = []
for k in range(1, 6):
    hdrs += [f"Produit{k}", f"Quantité{k}"]
hdrs += ["Suggested Package All Seasons", "Suggested Package Summer", "Info"]
for c, h in enumerate(hdrs, start=1):
    ws.cell(1, c, h)
style_hdr(ws, [f"{get_column_letter(c)}1" for c in range(1, len(hdrs) + 1)])
# lignes d'exemple : jusqu'à 5 produits (SKU) par commande
rows = [
    [("9004", 10), ("9130", 2), ("9124", 3)],
    [("9491", 5), ("9004", 1)],
    [("9124", 5), ("9130", 1)],
    [("9148", 1), ("9263", 2), ("243", 10), ("9055", 4), ("9138", 3)],
]
for i, prods in enumerate(rows, start=2):
    for k, (sku, q) in enumerate(prods):
        ws.cell(i, 2 * k + 1, sku); ws.cell(i, 2 * k + 2, q)
    for c in range(1, 11):
        ws.cell(i, c).border = box; ws.cell(i, c).font = normal
ws["P4"] = "Mode d'emploi : une commande par ligne, jusqu'à 5 produits (SKU ou nom exact) avec leur quantité (vide = 1)."
ws["P5"] = "Bouton « Remplir » : résultats en K, L (+ Info en M). Marges lues dans Packaging Calculator."
ws["P4"].font = Font(name="Arial", italic=True, color="666666"); ws["P5"].font = Font(name="Arial", italic=True, color="666666")
for k in range(1, 6):
    ws.column_dimensions[get_column_letter(2 * k - 1)].width = 12
    ws.column_dimensions[get_column_letter(2 * k)].width = 10
ws.column_dimensions["K"].width = 36; ws.column_dimensions["L"].width = 36; ws.column_dimensions["M"].width = 60
ws.column_dimensions["N"].width = 16
ws.freeze_panes = "A2"

# ---------- Référentiel Entrepôt (base de comparaison) ----------
import csv, os
ref_csv = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tests", "referentiel_entrepot.csv")
if os.path.exists(ref_csv):
    if "Référentiel Entrepôt" in wb.sheetnames: del wb["Référentiel Entrepôt"]
    ws = wb.create_sheet("Référentiel Entrepôt")
    rows = list(csv.DictReader(open(ref_csv, encoding="utf-8"), delimiter=";"))
    boxes = [k for k in rows[0].keys() if k not in ("categorie", "sku_candidats", "couchable")]
    grey = PatternFill("solid", fgColor="BFBFBF")
    # applique aux SKU candidats : Couchable? (si 'Oui' dans le CSV) et Caisses interdites (cases X), sans ecraser une valeur saisie
    hp = [str(c.value or "") for c in wsP[1]]
    cCouch = next(i for i, h in enumerate(hp, start=1) if h.lower().startswith("couchable"))
    cForb = next(i for i, h in enumerate(hp, start=1) if h.lower().startswith("caisses interdites"))
    rowsku = {str(wsP.cell(r, 1).value).strip(): r for r in range(2, wsP.max_row + 1) if wsP.cell(r, 1).value is not None}
    for r in rows:
        xboxes = [b for b in boxes if r[b].strip().upper() == "X"]
        for sku in [x.strip() for x in r["sku_candidats"].split("|") if x.strip()]:
            if sku not in rowsku: continue
            pr = rowsku[sku]
            if r.get("couchable", "").strip().lower() == "oui":
                wsP.cell(pr, cCouch).value = "Oui"
            if xboxes and not wsP.cell(pr, cForb).value:
                wsP.cell(pr, cForb).value = ";".join(xboxes)
    skuname = {str(wsP.cell(r, 1).value).strip(): wsP.cell(r, 3).value for r in range(2, wsP.max_row + 1) if wsP.cell(r, 1).value is not None}
    hdr = ["Catégorie entrepôt", "SKU candidats (séparés par |)", "Produit (1er SKU)"] + boxes + ["Couchable retenu", "Commentaire"]
    for c, h in enumerate(hdr, start=1): ws.cell(1, c, h)
    style_hdr(ws, [f"{get_column_letter(c)}1" for c in range(1, len(hdr) + 1)])
    for i, r in enumerate(rows, start=2):
        ws.cell(i, 1, r["categorie"]).font = bold
        ws.cell(i, 2, r["sku_candidats"]).fill = inp_fill
        first = r["sku_candidats"].split("|")[0]
        ws.cell(i, 3, skuname.get(first, "(à identifier)"))
        for j, b in enumerate(boxes, start=4):
            v = r[b].strip()
            cell = ws.cell(i, j)
            if v == "X":
                cell.value = "X"; cell.fill = grey
            elif v:
                cell.value = int(v) if v.isdigit() else v
            cell.alignment = Alignment(horizontal="center")
            if v != "X": cell.fill = inp_fill
        ws.cell(i, 4 + len(boxes), r.get("couchable", "Non"))
        for c in range(1, len(hdr) + 1): ws.cell(i, c).border = box
    n = len(rows) + 3
    ws.cell(n, 1, "Lecture : nombre MAX d'articles de la catégorie que la caisse peut contenir (« jusqu'à »). X (gris) = caisse non autorisée pour cet article.")
    ws.cell(n + 1, 1, "Cellules jaunes modifiables : SKU candidats (le 1er sert au calcul) et quantités. « >5 & <12 » = caisse utilisable seulement de 6 à 12 articles. Comparaison : tools/compare_entrepot.py ; recherche de réglages : tools/search_config.py.")
    ws.cell(n, 1).font = Font(name="Arial", italic=True, color="666666"); ws.cell(n + 1, 1).font = Font(name="Arial", italic=True, color="666666")
    ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 42; ws.column_dimensions["C"].width = 30
    for j in range(4, 4 + len(boxes)): ws.column_dimensions[get_column_letter(j)].width = 13
    ws.column_dimensions[get_column_letter(4 + len(boxes))].width = 16
    ws.column_dimensions[get_column_letter(5 + len(boxes))].width = 40
    ws.freeze_panes = "D2"

# police Arial sur les nouvelles feuilles (cellules non stylées)
for name in ("Packaging Calculator", "Mixed Pack Helper", "Référentiel Entrepôt"):
    if name not in wb.sheetnames: continue
    for row in wb[name].iter_rows():
        for c in row:
            if c.value is not None and c.font.name != "Arial":
                c.font = Font(name="Arial", bold=c.font.bold, italic=c.font.italic, color=c.font.color)
wb.save(dst)
print("saved", dst)
